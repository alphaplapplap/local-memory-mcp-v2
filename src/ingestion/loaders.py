#!/usr/bin/env python3
"""
Document loaders for various file formats.
"""

from pathlib import Path
from typing import Any, Dict, List

from .base import DocumentLoader


class TextLoader(DocumentLoader):
    """Loader for plain text files."""

    def can_handle(self, file_path: Path) -> bool:
        """Check if this loader can handle the given file."""
        return file_path.suffix.lower() in ['.txt', '.text']

    async def extract_chunks(self, file_path: Path, **kwargs):
        """Extract text chunks from a document."""
        from .chunker import TextChunker, ChunkingStrategy
        from .base import DocumentChunk
        
        # Load document content
        content = self.load_document(file_path)
        
        # Set up chunking strategy
        chunk_size = kwargs.get('chunk_size', 1000)
        chunk_overlap = kwargs.get('chunk_overlap', 200)
        
        strategy = ChunkingStrategy(
            chunk_size=chunk_size,
            chunk_overlap=chunk_overlap
        )
        chunker = TextChunker(strategy)
        
        # Create chunks
        chunks = chunker.chunk_text(content)
        
        # Yield DocumentChunk objects
        for i, (chunk_text, chunk_metadata) in enumerate(chunks):
            yield DocumentChunk(
                content=chunk_text,
                metadata=chunk_metadata,
                chunk_index=i,
                source_file=file_path
            )

    def load_document(self, file_path: Path) -> str:
        """Load text content from file."""
        try:
            return file_path.read_text(encoding="utf-8")
        except UnicodeDecodeError:
            # Try with different encodings
            for encoding in ["latin-1", "cp1252", "iso-8859-1"]:
                try:
                    return file_path.read_text(encoding=encoding)
                except UnicodeDecodeError:
                    continue
            raise ValueError(
                f"Could not decode file {file_path} with any supported encoding"
            )


class MarkdownLoader(DocumentLoader):
    """Loader for Markdown files."""

    def can_handle(self, file_path: Path) -> bool:
        """Check if this loader can handle the given file."""
        return file_path.suffix.lower() in ['.md', '.markdown']

    async def extract_chunks(self, file_path: Path, **kwargs):
        """Extract text chunks from a markdown document."""
        from .chunker import TextChunker, ChunkingStrategy
        from .base import DocumentChunk
        
        # Load document content
        content = self.load_document(file_path)
        
        # Set up chunking strategy
        chunk_size = kwargs.get('chunk_size', 1000)
        chunk_overlap = kwargs.get('chunk_overlap', 200)
        
        strategy = ChunkingStrategy(
            chunk_size=chunk_size,
            chunk_overlap=chunk_overlap
        )
        chunker = TextChunker(strategy)
        
        # Create chunks
        chunks = chunker.chunk_text(content)
        
        # Yield DocumentChunk objects
        for i, (chunk_text, chunk_metadata) in enumerate(chunks):
            yield DocumentChunk(
                content=chunk_text,
                metadata=chunk_metadata,
                chunk_index=i,
                source_file=file_path
            )

    def load_document(self, file_path: Path) -> str:
        """Load Markdown content from file."""
        return file_path.read_text(encoding="utf-8")


# Optional loaders that require additional dependencies
try:
    import PyPDF2

    class PDFLoader(DocumentLoader):
        """Loader for PDF files using PyPDF2."""

        def load_document(self, file_path: Path) -> str:
            """Extract text from PDF file."""
            with open(file_path, "rb") as file:
                pdf_reader = PyPDF2.PdfReader(file)
                text = ""
                for page in pdf_reader.pages:
                    text += page.extract_text() + "\n"
                return text

except ImportError:
    # PyPDF2 not available
    pass


try:
    from docx import Document

    class DOCXLoader(DocumentLoader):
        """Loader for DOCX files using python-docx."""

        def load_document(self, file_path: Path) -> str:
            """Extract text from DOCX file."""
            doc = Document(file_path)
            text = ""
            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"
            return text

except ImportError:
    # python-docx not available
    pass


def load_document(file_path: str) -> str:
    """Convenience function to load a document.
    
    Args:
        file_path: Path to the document file
        
    Returns:
        The document content as text
    """
    path = Path(file_path)
    
    # Simple file reading without using loader classes
    # This avoids the abstract class instantiation issues
    suffix = path.suffix.lower()
    
    try:
        if suffix == '.txt' or suffix == '':
            # Plain text file
            return path.read_text(encoding="utf-8")
        elif suffix in ['.md', '.markdown']:
            # Markdown file
            return path.read_text(encoding="utf-8")
        elif suffix == '.pdf':
            # Try PDF loading if PyPDF2 is available
            try:
                import PyPDF2
                with open(path, "rb") as file:
                    pdf_reader = PyPDF2.PdfReader(file)
                    text = ""
                    for page in pdf_reader.pages:
                        text += page.extract_text() + "\n"
                    return text
            except ImportError:
                raise ImportError("PyPDF2 not available for PDF loading")
        elif suffix == '.docx':
            # Try DOCX loading if python-docx is available
            try:
                from docx import Document
                doc = Document(path)
                text = ""
                for paragraph in doc.paragraphs:
                    text += paragraph.text + "\n"
                return text
            except ImportError:
                raise ImportError("python-docx not available for DOCX loading")
        else:
            # Default to text file
            return path.read_text(encoding="utf-8")
            
    except UnicodeDecodeError:
        # Try with different encodings
        for encoding in ["latin-1", "cp1252", "iso-8859-1"]:
            try:
                return path.read_text(encoding=encoding)
            except UnicodeDecodeError:
                continue
        raise ValueError(f"Could not decode file {path} with any supported encoding")
